from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts"
DOCX = OUTPUT / "Node_Solutions_AI_Triage_Technical_Report.docx"

NAVY = "17365D"
TEAL = "0F766E"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F5F2"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"
WHITE = "FFFFFF"
BLACK = RGBColor(0, 0, 0)


def shade(cell, fill: str) -> None:
    """Apply deterministic table-cell shading."""
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def borders(table) -> None:
    """Set visible light-gray borders on every table edge."""
    props = table._tbl.tblPr
    element = props.find(qn("w:tblBorders"))
    if element is None:
        element = OxmlElement("w:tblBorders")
        props.append(element)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), BORDER)
        element.append(edge)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    """Give table text enough whitespace to remain readable after PDF rendering."""
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table(doc, headers: list[str], rows: list[list[str]], widths=None):
    """Create a consistently styled comparison or mapping table."""
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    borders(result)
    header_props = result.rows[0]._tr.get_or_add_trPr()
    header_props.append(OxmlElement("w:tblHeader"))
    for index, heading in enumerate(headers):
        cell = result.rows[0].cells[index]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(heading)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = result.add_row().cells
        for index, value in enumerate(values):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cells[index], PALE_BLUE if row_index % 2 else WHITE)
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(9)
    for row in result.rows:
        for index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return result


def bullet(doc, text: str, level=0) -> None:
    """Add a readable bullet with an optional nesting level."""
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(text, style=style)
    paragraph.paragraph_format.space_after = Pt(3)


def numbered(doc, text: str) -> None:
    """Add one ordered process step."""
    paragraph = doc.add_paragraph(text, style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)


def code(doc, text: str) -> None:
    """Render a compact code or command block without a decorative callout box."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    for index, line in enumerate(text.splitlines()):
        run = paragraph.add_run(line)
        run.font.name = "Aptos Mono"
        run.font.size = Pt(8.5)
        if index < len(text.splitlines()) - 1:
            run.add_break()


def heading(doc, text: str, level: int = 1) -> None:
    """Add a heading that remains with its first following paragraph."""
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def paragraph(doc, text: str, bold_lead: str | None = None) -> None:
    """Add normal report prose with an optional bold opening label."""
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(7)


def page_title(doc, number: str, title: str, intro: str) -> None:
    """Start a major report chapter on a clean new page."""
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(number.upper())
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(15, 118, 110)
    heading(doc, title, 1)
    paragraph(doc, intro)


def flow_table(doc) -> None:
    """Build an editable left-to-right architecture flow using native Word tables."""
    items = [
        ("01", "Streamlit", "Input and review"),
        ("02", "FastAPI", "Validate and control"),
        ("03", "LangGraph", "Four typed stages"),
        ("04", "Ollama", "Local AI roles"),
        ("05", "PostgreSQL", "Decision and review"),
    ]
    result = doc.add_table(rows=1, cols=len(items))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    header_props = result.rows[0]._tr.get_or_add_trPr()
    header_props.append(OxmlElement("w:tblHeader"))
    for index, (number, title, detail) in enumerate(items):
        cell = result.rows[0].cells[index]
        shade(cell, PALE_TEAL if index % 2 == 0 else PALE_BLUE)
        set_cell_margins(cell, 150, 100, 150, 100)
        cell.width = Inches(1.38)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(number + "\n")
        r.bold = True
        r.font.color.rgb = RGBColor(15, 118, 110)
        r = p.add_run(title + "\n")
        r.bold = True
        r.font.size = Pt(9.5)
        r = p.add_run(detail)
        r.font.size = Pt(8)
    borders(result)
    doc.add_paragraph()


def build() -> None:
    """Create the complete assessment report as a styled Word document."""
    OUTPUT.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = BLACK
    styles["Normal"].paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 30), ("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 11.5)):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = BLACK
        styles[name].paragraph_format.space_before = Pt(10)
        styles[name].paragraph_format.space_after = Pt(6)
    styles["Title"].paragraph_format.space_after = Pt(12)
    # Some renderers add a theme border to Word's built-in Title style; remove it explicitly.
    title_properties = styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(50)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("AI Request Triage Assistant")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Technical design implementation and evaluation report").bold = True
    subtitle.runs[0].font.size = Pt(16)
    subtitle.paragraph_format.space_after = Pt(20)
    paragraph(
        doc,
        "A local, review-first agent workflow for professional-services request classification, priority "
        "reasoning, ownership routing, and safe customer-response drafting.",
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(90)
    table(doc, ["Submission", "Implementation"], [
        ["Candidate", "Akshita Rastogi"],
        ["Challenge", "Node Solutions AI Technical Challenge"],
        ["Runtime", "Docker, FastAPI, Streamlit, PostgreSQL, Redis"],
        ["AI", "Local Ollama with two schema-constrained roles"],
        ["Orchestration", "LangGraph typed state workflow"],
    ], [2.0, 4.7])
    paragraph(doc, "The system supports the six supplied mock requests and new mock requests. Every output "
                   "is a draft for human review; no endpoint sends messages or changes a customer system.")

    # Executive summary
    page_title(doc, "01", "Executive summary", "The implementation satisfies the assessment while adding "
               "targeted controls for the failure modes that matter most in real AI workflows.")
    paragraph(doc, "The assistant accepts a written request and returns a short summary, exactly one category, "
                   "one priority with a reason, one allowed owner, and a professional first-response draft. "
                   "The six supplied messages are data fixtures rather than hardcoded response branches; the "
                   "same endpoint accepts any validated mock request between 5 and 5,000 characters.")
    paragraph(doc, "The main engineering decision is to separate model judgment from business safety. Ollama "
                   "handles summarisation and nuanced classification through a strict JSON schema. LangGraph "
                   "then applies deterministic policy before drafting. Privacy exposure, outages, billing, "
                   "commercial intent, and product enhancements therefore have stable routing invariants.")
    heading(doc, "Results at a glance", 2)
    table(doc, ["Dimension", "Implemented approach", "Why it matters"], [
        ["Local AI", "Ollama model on the host", "No hosted inference key or cloud data transfer"],
        ["Reliability", "Schema validation, repair, and fallback", "Malformed output cannot break the contract"],
        ["Safety", "Redaction, injection branch, policy rules", "High-cost cases do not rely only on prompting"],
        ["Review", "Immutable decision plus separate review", "AI history is never overwritten"],
        ["Operations", "Idempotency, rate limits, health, metrics", "Retries and failures are visible"],
        ["Evaluation", "Golden cases plus adversarial tests", "Routing regressions are reproducible"],
    ], [1.1, 2.4, 3.2])

    # Requirements
    page_title(doc, "02", "Assessment requirements and coverage", "Each requested capability maps to an "
               "observable implementation and a concrete demonstration path.")
    table(doc, ["Assessment requirement", "Implementation", "Demonstration evidence"], [
        ["Written request", "Streamlit form and POST /v1/triage", "Use any supplied or custom mock text"],
        ["Short summary", "AnalysisDecision.summary", "Visible above category metrics"],
        ["Five categories", "Strict Category enum", "Invalid values fail schema validation"],
        ["Four priorities", "Strict Priority enum plus rationale", "Priority and reason appear together"],
        ["Allowed owners", "Strict Owner enum and policy rules", "Owner shown in UI and database"],
        ["First response", "Separate response role or guarded template", "Always labelled for review"],
        ["Six requests", "data/mock_requests.json", "All six appear in the selector"],
        ["New requests", "Generic typed endpoint", "Six extra edge cases plus custom input"],
        ["Request 05", "Privacy incident invariant", "Urgent Engineering route and risk flags"],
        ["Decisions and limitations", "README and report", "Trade-offs are stated explicitly"],
    ], [1.65, 2.55, 2.55])
    heading(doc, "Expected routing for the supplied requests", 2)
    table(doc, ["ID", "Primary intent", "Category", "Priority", "Owner"], [
        ["01", "Automation discovery call", "Sales", "Medium", "Sales Team"],
        ["02", "Active portal outage", "Technical", "Urgent", "Engineering"],
        ["03", "Duplicate invoice charge", "Billing", "High", "Finance"],
        ["04", "Future product enhancement", "Technical", "Low", "Engineering"],
        ["05", "Wrong-workspace data exposure", "Technical", "Urgent", "Engineering"],
        ["06", "Custom AI system enquiry", "Sales", "Medium", "Sales Team"],
    ], [0.4, 2.2, 1.1, 0.9, 1.35])

    # Architecture
    page_title(doc, "03", "Architecture", "The design uses a small number of components with explicit "
               "responsibilities and stable boundaries.")
    flow_table(doc)
    paragraph(doc, "The browser talks only to FastAPI. The API authenticates, validates, rate-limits, and checks "
                   "idempotency before invoking the graph. LangGraph passes typed state through sanitation, "
                   "analysis, validation, and drafting. PostgreSQL stores the final decision and human review; "
                   "Redis stores only short-lived rate counters. Structured logs and Prometheus metrics observe "
                   "the HTTP and decision paths without using request text as labels.")
    heading(doc, "Component choices", 2)
    table(doc, ["Component", "Responsibility", "Decision rationale"], [
        ["Streamlit", "Operator input, results, trace, review queue", "Fast local UI with no production UI requirement"],
        ["FastAPI", "Typed HTTP boundary and status semantics", "Async, OpenAPI-native, integrates with Pydantic"],
        ["LangGraph", "Explicit state and ordered node execution", "Control flow and failure paths remain visible"],
        ["Ollama", "Local analysis and response generation", "Private local inference and structured output"],
        ["PostgreSQL", "Durable decisions and human reviews", "Uniqueness, foreign keys, ordering, and audit integrity"],
        ["Redis", "Shared rate-limit counters", "Expiry and future multi-replica consistency"],
    ], [1.15, 2.3, 3.0])
    heading(doc, "Why there is no vector database", 2)
    paragraph(doc, "This challenge is classification and drafting, not knowledge retrieval. There is no policy "
                   "corpus to embed or cite. Adding Qdrant would create operational surface without improving "
                   "the requested result. If policy documents are introduced, a retrieval node can be inserted "
                   "before analysis and carry excerpts and citations in graph state.")

    # LangGraph
    page_title(doc, "04", "LangGraph workflow and state", "The graph is an auditable state machine, not an "
               "open-ended autonomous agent.")
    table(doc, ["Node", "Input", "Output", "Failure behavior"], [
        ["Sanitize", "Original request", "Redacted text, injection flag", "Reject invalid API input before graph"],
        ["Analyse", "Redacted text", "Typed AnalysisDecision", "One repair, then deterministic fallback"],
        ["Validate", "Decision and request signals", "Policy-corrected decision", "Non-negotiable invariants win"],
        ["Draft", "PII-free summary and decision", "Customer-facing draft", "Repair, template, or guarded branch"],
    ], [1.0, 1.55, 2.1, 2.0])
    heading(doc, "State fields", 2)
    table(doc, ["Field", "Producer", "Purpose"], [
        ["original_text", "Service entry", "Faithful controlled audit input"],
        ["redacted_text", "Sanitize", "Only request form eligible for analysis inference"],
        ["injection_detected", "Sanitize", "Conditional safety decision"],
        ["decision", "Analyse and Validate", "Schema-safe classification and routing"],
        ["decision_source", "Analyse or injection policy", "Distinguishes model, fallback, and guardrail"],
        ["warnings", "All relevant stages", "Add why the path changed or degraded"],
        ["trace", "Every graph node", "Attempts, duration, tokens, corrections, and branches"],
        ["draft", "Response agent or template", "Review-only first response"],
    ], [1.35, 1.45, 3.9])
    paragraph(doc, "Warnings and trace use additive reducers. Each node appends its evidence rather than "
                   "overwriting prior stages. The final UI trace therefore explains the complete route.")
    heading(doc, "Conditional safety branches", 2)
    bullet(doc, "Prompt injection: normalise to Other, Low, Client Success and use a guarded clarification template.")
    bullet(doc, "Privacy or access-control incident: preserve Urgent Engineering routing and use a deterministic incident acknowledgement that never claims remediation occurred.")
    bullet(doc, "Normal request: use the response agent with a PII-free validated summary.")

    # Framework decision
    page_title(doc, "05", "Why LangGraph instead of other frameworks", "The selected framework matches a "
               "known, safety-ordered workflow and leaves room for controlled expansion.")
    table(doc, ["Option", "Strength", "Reason for this decision"], [
        ["LangGraph", "Typed state, reducers, conditional edges, interrupts", "Selected: control flow and evidence are explicit"],
        ["LangChain", "Model, prompt, retrieval, and tool composition", "Useful ecosystem, but a linear chain hides state transitions"],
        ["ReAct agent", "Dynamic tool and step selection", "Rejected: unnecessary autonomy and variable latency"],
        ["CrewAI or AutoGen", "Conversational multi-agent teams", "Rejected: two ordered roles do not need agent messaging"],
        ["Plain Python", "Minimal dependencies and direct control", "Viable now, but future branches and interrupts are less visible"],
    ], [1.35, 2.5, 2.9])
    paragraph(doc, "The workflow is sequential because each stage depends on the prior result: sanitation must "
                   "precede inference, validation must follow analysis, and drafting must use the validated "
                   "decision. Parallel agents would add coordination cost without shortening this critical path.")
    heading(doc, "Why two AI roles", 2)
    bullet(doc, "The analysis role uses temperature 0 and a Pydantic JSON schema for stable routing fields.")
    bullet(doc, "The response role uses temperature 0.1 and a DraftOutput schema for constrained language.")
    bullet(doc, "Separating them prevents customer-facing tone generation from changing category, priority, or owner.")

    # End to end
    page_title(doc, "06", "End to end code execution", "This section follows one request from the browser to "
               "durable storage and human review.")
    steps = [
        "Docker starts uvicorn triage.main:app. FastAPI lifespan opens PostgreSQL and Redis clients, creates the Ollama client, compiles the graph once, and constructs TriageService.",
        "Streamlit api() sends POST /v1/triage with request text, source, API key, and a unique idempotency key.",
        "request_context() binds a correlation ID, measures latency, records status metrics, and writes a structured completion log.",
        "TriageRequest validates length, source, meaningful content, and normalised whitespace. require_api_key() authenticates without logging the key.",
        "RateLimiter.allow() increments a Redis minute bucket. TriageService checks PostgreSQL for a prior successful result with the same idempotency key.",
        "The sanitize node redacts contact and payment identifiers, detects injection language, and appends warnings and trace evidence.",
        "The analysis node calls OllamaClient.analyse(). Ollama receives the AnalysisDecision JSON schema; Pydantic validates the returned values. Invalid output gets one repair attempt.",
        "The validation node applies incident, outage, billing, feature, and commercial ownership invariants. It records any correction visibly.",
        "The drafting stage sees only the validated summary and decision. High-risk and injection branches use deterministic templates; normal requests use structured local generation.",
        "Database.insert() writes one complete result. The API returns HTTP 201, or HTTP 200 with Idempotent-Replayed for a replay.",
        "show_result() renders the decision and trace. A human review is written to triage_reviews without changing the original inference.",
    ]
    for item in steps:
        numbered(doc, item)

    # Files
    page_title(doc, "07", "Code structure and method guide", "Each module has one primary reason to change.")
    table(doc, ["File", "Main symbols", "Role"], [
        ["main.py", "lifespan, middleware, routes", "HTTP boundary, lifecycle, metrics, status codes"],
        ["models.py", "enums, request, decision, result", "One source of truth for typed contracts"],
        ["graph.py", "TriageState, build_graph", "Nodes, edges, conditional safety paths"],
        ["ollama.py", "health, analyse, draft", "Local inference, schemas, timeout, repair"],
        ["guardrails.py", "safe_fallback, enforce_business_rules", "Model-independent routing policy"],
        ["security.py", "redaction, injection, authentication", "Untrusted-input boundary"],
        ["service.py", "TriageService.triage", "Idempotency, graph invocation, persistence"],
        ["database.py", "connect, insert, get, list, review", "PostgreSQL schema and repository"],
        ["rate_limit.py", "allow, healthy, close", "Redis fixed-window operational control"],
        ["metrics.py", "REQUESTS, LATENCY, DECISIONS", "Low-cardinality Prometheus telemetry"],
        ["ui/app.py", "api, show_result", "Input, history, queue, evaluation, explanation"],
        ["evaluate.py", "main", "Golden set through the real HTTP API"],
    ], [1.25, 2.25, 3.2])
    heading(doc, "Adding new behavior safely", 2)
    paragraph(doc, "A new request needs no ingestion: paste mock text or call the endpoint. To add a demo fixture, "
                   "append an object to edge_case_requests.json. To add a category, update the enum, prompt, policy "
                   "rules, UI expectations, golden cases, and tests together. To add a tool, define typed input and "
                   "output, a conditional graph edge, timeout and retry limits, and persisted trace evidence.")

    # Guardrails
    page_title(doc, "08", "Reliability privacy and guardrails", "Controls are layered because no single prompt "
               "or confidence value is a sufficient safety boundary.")
    table(doc, ["Risk", "Control", "Visible evidence"], [
        ["Malformed model JSON", "Pydantic schema plus one repair", "Attempts and fallback warning"],
        ["Model unavailable", "Deterministic decision and draft", "decision_source and trace status"],
        ["Prompt injection", "Detection and conditional isolation", "prompt_injection flag and guarded template"],
        ["PII leakage", "Pre-inference redaction and summary-only drafting", "Redaction types in warnings and trace"],
        ["Under-prioritised incident", "Security and outage invariants", "Policy correction warning"],
        ["False remediation claim", "Security response template and draft validator", "No completion claim reaches UI"],
        ["Duplicate client retry", "PostgreSQL idempotency uniqueness", "HTTP 200 replay header"],
        ["Abuse or burst", "Redis shared fixed-window limit", "HTTP 429 and expiring counter"],
        ["Hidden dependency failure", "Readiness probes", "HTTP 503 with per-dependency state"],
    ], [1.55, 2.55, 2.6])
    heading(doc, "Draft-specific validation", 2)
    bullet(doc, "Reject reasoning markers such as 'the user is asking' or 'I need to'.")
    bullet(doc, "Reject redaction placeholders and instructions to contact a placeholder.")
    bullet(doc, "Reject unverified claims such as 'we have confirmed', 'we initiated', or 'we removed'.")
    bullet(doc, "Reject Urgent or immediate language when the validated priority is not Urgent.")
    bullet(doc, "Require a complete 20-800 character structured message and terminal punctuation.")

    # Edge cases
    page_title(doc, "09", "Additional request coverage", "Six extra cases demonstrate generalisation and "
               "adversarial handling beyond the supplied fixtures.")
    table(doc, ["Case", "What it tests", "Expected behavior"], [
        ["E01 Ambiguous", "Insufficient specialist evidence", "Other and human clarification"],
        ["E02 Injection", "Embedded prompt instructions", "Isolate, flag, and guarded response"],
        ["E03 Identifiers", "Email and phone plus billing", "Redact, retain Billing High Finance"],
        ["E04 Conflict", "Customer says not urgent but portal is down", "Impact wins: Technical Urgent Engineering"],
        ["E05 Mixed intent", "Billing and sales in one message", "Choose one primary operational category"],
        ["E06 Novel support", "Unseen account-administration question", "Support and Client Success"],
    ], [1.25, 2.65, 2.8])
    paragraph(doc, "These examples are not proof of universal coverage. Novel languages, subtle legal requests, "
                   "and unseen security paraphrases require a larger reviewed dataset. Human review remains the "
                   "final safety boundary.")

    # Persistence
    page_title(doc, "10", "Persistence observability and operations", "The local deployment can be inspected "
               "through both terminal commands and desktop database tools.")
    heading(doc, "PostgreSQL records", 2)
    table(doc, ["Table", "Stored information", "Integrity"], [
        ["triage_requests", "Input, redacted form, decision, draft, warnings, trace", "UUID primary key and unique idempotency key"],
        ["triage_reviews", "Outcome, notes, review timestamp", "One review per request with foreign key"],
    ], [1.45, 3.2, 2.1])
    paragraph(doc, "DBeaver connection: host localhost, port 5433, database triage, username triage, password "
                   "triage. RedisInsight connection: host localhost, port 6380, no username or password. Redis "
                   "rate keys expire after 60 seconds, so refresh immediately after an API request.")
    heading(doc, "Operational commands", 2)
    code(doc, "curl -s http://localhost:8100/health/ready | python -m json.tool\n"
              "docker compose logs --tail=100 api\n"
              "curl -s http://localhost:8100/metrics | grep '^triage_'\n"
              "docker compose exec redis redis-cli --scan --pattern 'rate:*'")
    heading(doc, "HTTP semantics", 2)
    table(doc, ["Status", "Meaning"], [
        ["200", "Read operation or idempotent replay"], ["201", "New triage result or review"],
        ["401", "Missing or invalid API key"], ["404", "Unknown request ID"],
        ["409", "Duplicate review or idempotency race"], ["422", "Invalid body, enum, UUID, or limit"],
        ["429", "Rate limit exceeded"], ["503", "Required dependency is unavailable"],
        ["500", "Unexpected internal failure with safe public detail"],
    ], [1.0, 5.7])

    # Evaluation
    page_title(doc, "11", "Evaluation strategy", "Evaluation checks deterministic routing today and defines "
               "the measurements needed before a real rollout.")
    paragraph(doc, "data/evaluation.json contains the six supplied requests with expected category, priority, "
                   "and owner. scripts/evaluate.py calls the real HTTP boundary, compares all three fields, and "
                   "writes detailed expected-versus-actual evidence to output/evaluation-results.json.")
    table(doc, ["Metric", "Purpose", "Release interpretation"], [
        ["Category macro F1", "Balance performance across categories", "Avoid majority-class accuracy"],
        ["Owner exact match", "Verify operational routing", "Wrong team is a workflow failure"],
        ["Critical false-negative rate", "Measure missed outages/privacy cases", "Release-blocking safety metric"],
        ["Schema-valid rate", "Track structured-output reliability", "Should remain near 100 percent"],
        ["Human edit distance", "Approximate draft usefulness", "Lower is better after reviewer agreement"],
        ["p50 p95 p99 latency", "Expose slow-tail behavior", "Regression gate by environment"],
        ["Fallback rate", "Detect model or prompt degradation", "Investigate sustained increases"],
    ], [1.55, 2.9, 2.25])
    heading(doc, "Commands", 2)
    code(doc, "docker compose exec api pytest -q\n"
              "docker compose exec api ruff check --line-length 100 src tests scripts ui\n"
              "docker compose exec api python scripts/evaluate.py\n"
              "python -m json.tool output/evaluation-results.json")
    paragraph(doc, "The current model confidence is diagnostic metadata, not a calibrated probability and not "
                   "permission to auto-send. A larger evaluation set would be versioned, multilingual, balanced, "
                   "paraphrase-heavy, and independently labelled by multiple reviewers.")

    # Brownie points
    page_title(doc, "12", "Additional engineering features", "The extensions focus on trustworthy operation, "
               "not decorative feature count.")
    extras = [
        ("Typed agent state", "Separate roles and additive trace reducers make every transition auditable."),
        ("Structured local output", "Both Ollama roles return Pydantic-validated JSON with bounded repair."),
        ("Deterministic policy", "High-cost routing is enforced outside probabilistic model behavior."),
        ("Data minimisation", "The drafting role sees only a validated PII-free summary and decision."),
        ("Conditional guardrails", "Injection and privacy paths bypass normal generation when autonomy adds risk."),
        ("Graceful degradation", "Local model failure returns safe, visible fallback behavior."),
        ("Immutable review", "Human feedback is captured without rewriting the historical model output."),
        ("Safe retry semantics", "Idempotency and database uniqueness prevent duplicate successful work."),
        ("Operational visibility", "Correlation logs, dependency checks, metrics, and graph traces aid diagnosis."),
        ("Adversarial evaluation", "Tests intentionally challenge routing, PII, injection, and reasoning leakage."),
    ]
    table(doc, ["Feature", "Engineering value"], [[a, b] for a, b in extras], [1.8, 4.9])

    # Limitations
    page_title(doc, "13", "Limitations and next improvements", "The project is production-aware but does not "
               "claim that a local assessment application is production-ready.")
    heading(doc, "Known limitations", 2)
    bullet(doc, "The deterministic vocabulary is English and finite; new paraphrases need evaluated policy updates.")
    bullet(doc, "Six golden examples cannot estimate real generalisation, class imbalance, or multilingual quality.")
    bullet(doc, "The demo uses one API-key role and has no tenant row-level security or secrets manager.")
    bullet(doc, "Model confidence is not calibrated against held-out reviewer labels.")
    bullet(doc, "The fixed-window limiter can permit a boundary burst, although it is shared across replicas.")
    bullet(doc, "No endpoint sends email, edits workspace access, or performs remediation in customer systems.")
    heading(doc, "Next improvements", 2)
    paragraph(doc, "Expand the versioned dataset with multilingual, ambiguous, conflicting, and reviewer-disagreement cases.", "1. ")
    paragraph(doc, "Calibrate critical-risk scores and add release gates for false negatives, schema validity, and latency percentiles.", "2. ")
    paragraph(doc, "Add role-based access, tenant isolation, managed secrets, migrations, and encrypted audit retention policies.", "3. ")
    paragraph(doc, "Introduce OpenTelemetry spans and dashboards for per-node latency, model fallback rate, and reviewer corrections.", "4. ")
    paragraph(doc, "Use reviewed corrections as offline evaluation or supervised-learning data only after governance approval.", "5. ")
    paragraph(doc, "Capturing approval or rejection is not reinforcement learning or RLHF by itself. Feedback becomes "
                   "useful only after curation, quality review, and a separate offline improvement process.")

    # Demonstration
    page_title(doc, "14", "Demonstration guide", "A short demonstration should show correctness, explainability, "
               "a high-risk branch, persistence, and evaluation rather than every screen control.")
    table(doc, ["Time", "Screen", "What to demonstrate"], [
        ["0:00-0:40", "README and architecture", "Problem, local stack, four graph stages, and review-first boundary"],
        ["0:40-1:30", "Request 01", "Sales correction, structured draft, node attempts, tokens, and latency"],
        ["1:30-2:15", "Request 02", "Impact-driven Urgent Engineering outage routing"],
        ["2:15-3:10", "Request 05", "Privacy flags and deterministic incident response without false remediation"],
        ["3:10-3:55", "Edge E03", "Email and phone redaction plus Billing High Finance decision"],
        ["3:55-4:30", "Edge E02", "Conditional injection isolation and guarded clarification template"],
        ["4:30-5:15", "Review queue and DBeaver", "Immutable decision and separate human review rows"],
        ["5:15-5:50", "RedisInsight and metrics", "Expiring hashed rate key and low-cardinality telemetry"],
        ["5:50-6:30", "Evaluation", "Golden expected-versus-actual checks and honest limitations"],
    ], [0.9, 1.7, 4.1])
    heading(doc, "Concise narration", 2)
    paragraph(doc, "I built a local, review-first request triage assistant. FastAPI validates the request and "
                   "applies API-key, rate-limit, and idempotency controls. LangGraph then sanitises input, asks "
                   "Ollama for a schema-constrained decision, enforces deterministic routing policy, and creates "
                   "a reviewed draft. PostgreSQL preserves both the AI result and a separate human decision. "
                   "The most important addition is defense in depth: privacy incidents, outages, prompt injection, "
                   "PII, and false remediation claims are handled outside model optimism. The evaluation harness "
                   "runs all six supplied cases through the real API and records field-level evidence.")

    # Runbook
    page_title(doc, "15", "Local runbook", "The Docker ports are intentionally separate from the SupplyGuard "
               "project so both applications can coexist.")
    heading(doc, "Start", 2)
    code(doc, "cd /Users/akshitarastogi/Documents/Codex/2026-09-09/si/outputs/node-solutions-ai-triage\n"
              "ollama pull qwen3:4b\n"
              "cp .env.example .env\n"
              "docker compose up -d --build\n"
              "docker compose ps")
    table(doc, ["Surface", "Address"], [
        ["Streamlit", "http://localhost:8601"],
        ["FastAPI documentation", "http://localhost:8100/docs"],
        ["Prometheus metrics", "http://localhost:8100/metrics"],
        ["PostgreSQL", "localhost:5433 / database triage"],
        ["Redis", "localhost:6380"],
        ["Ollama host", "localhost:11434"],
    ], [2.25, 4.45])
    heading(doc, "Inspect and stop", 2)
    code(doc, "curl -s http://localhost:8100/health/ready | python -m json.tool\n"
              "docker compose logs -f api\n"
              "docker compose down")
    paragraph(doc, "docker compose down preserves the named PostgreSQL and Redis volumes. Use down -v only for "
                   "an intentional destructive reset.")

    # Conclusion
    page_title(doc, "16", "Conclusion", "The solution meets the challenge with a deliberately bounded and "
               "explainable agent system.")
    paragraph(doc, "The assistant is more than a classification prompt: it is a typed workflow with local "
                   "inference, deterministic safety policy, controlled degradation, durable audit history, and "
                   "human review. Additional complexity is justified by observable failure modes rather than by "
                   "an attempt to imitate a large production platform.")
    paragraph(doc, "The design can evolve without changing its central contract. Retrieval can be added when "
                   "policy documents exist; specialist tools can be added behind typed conditional edges; and "
                   "reviewer corrections can expand a governed offline evaluation set. Until then, the current "
                   "boundaries keep the assessment simple enough to explain and strong enough to demonstrate "
                   "engineering judgment.")

    # Consistent footer and compact header on content pages.
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "AI Request Triage Assistant   Technical Report"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 90, 90)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Node Solutions AI Technical Challenge")
        for run in footer.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 90, 90)

    doc.core_properties.title = "AI Request Triage Assistant Technical Report"
    doc.core_properties.subject = "Node Solutions AI Technical Challenge"
    doc.core_properties.author = "Akshita Rastogi"
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
